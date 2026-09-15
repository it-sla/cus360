from io import BytesIO
import uuid
from openpyxl import Workbook
from sqlalchemy import select
from app.company_imports import analyze,import_companies,parse_workbook
from app.db import SessionLocal
from app.models import Company,CompanyAlias,CompanyImportBatch,DataQualityIssue
from app.utils import normalize_icris

def workbook(rows,header=('Account Number','Company Name ',None)):
    book=Workbook();base=book.active;base.title=' BASE SHEET ';base.append(header)
    for row in rows:base.append((*row,None) if len(row)==2 else row)
    pivot=book.create_sheet('PIVOT TABLE');pivot.append(('Account Number','Company Name'));pivot.append(('DO-NOT-IMPORT','Pivot Co'))
    output=BytesIO();book.save(output);return output.getvalue()

def unique(prefix):return f'{prefix}-{uuid.uuid4().hex[:10]}'

def test_headers_blank_column_pivot_and_normalization():
    token=unique('HDR');data=workbook([(f'\u00a0 {token.lower()} \u00a0','Cargo One')])
    parsed=parse_workbook(data,'master.xlsx');assert parsed['worksheet_name']=='BASE SHEET';assert parsed['headers']==['Account Number','Company Name'];assert parsed['available_worksheets']==['BASE SHEET'];assert parsed['rows'][0]['normalized_icris']==token.upper()
    assert normalize_icris('   r1124y')=='R1124Y'

def test_identifier_formatted_as_text_preserves_zeroes():
    book=Workbook();ws=book.active;ws.title='BASE SHEET';ws.append(('Account Number','Company Name'))
    cell=ws.cell(2,1,123);cell.number_format='00000';ws.cell(2,2,'Leading Zero Cargo');out=BytesIO();book.save(out)
    assert parse_workbook(out.getvalue(),'master.xlsx')['rows'][0]['normalized_icris']=='00123'

def test_preview_groups_duplicate_and_conflicting_names():
    token=unique('DUP');db=SessionLocal()
    try:
        preview=analyze(db,parse_workbook(workbook([(token,'Name One'),(token.lower(),'Name Two')]),'master.xlsx'))
        assert preview['total_source_rows']==2;assert preview['unique_normalized_icris_count']==1;assert preview['duplicate_icris_group_count']==1;assert preview['name_conflict_count']==1
    finally:db.close()

def test_import_promotes_provisional_preserves_alias_and_is_idempotent():
    token=unique('PROMOTE');db=SessionLocal()
    try:
        company=Company(icris_number=token,company_name='CRM Source Name',normalized_name='CRM SOURCE NAME',source='crm',name_source='crm',is_provisional=True);db.add(company);db.flush()
        # A crm_icris_not_in_master issue raised for this provisional company must resolve
        # automatically the moment the master import promotes it — this is the one place
        # that issue type is ever meant to close on its own.
        open_issue=DataQualityIssue(issue_type='crm_icris_not_in_master',severity='info',status='open',company_id=company.id,source_company_name='CRM Source Name',details_json={});db.add(open_issue);db.commit()
        data=workbook([(token.lower(),'Official Master Name')]);first=import_companies(db,data,'master.xlsx','BASE SHEET',{'ICRIS Number':'Account Number','Company Name':'Company Name'});db.commit()
        db.refresh(company);assert not company.is_provisional;assert company.company_name=='Official Master Name';assert first.promoted_count==1;assert db.scalar(select(CompanyAlias).where(CompanyAlias.company_id==company.id,CompanyAlias.alias_name=='CRM Source Name'))
        db.refresh(open_issue);assert open_issue.status=='resolved';assert open_issue.resolved_by=='company_master_import'
        second=import_companies(db,data,'master.xlsx','BASE SHEET',{'ICRIS Number':'Account Number','Company Name':'Company Name'});db.commit();assert second.created_count==0;assert second.promoted_count==0;assert second.alias_added_count==0
    finally:db.close()

def test_repeat_conflict_import_reuses_quality_issue():
    token=unique('ISSUE');data=workbook([(token,'First Name'),(token,'Second Name')]);db=SessionLocal()
    try:
        for _ in range(2):import_companies(db,data,'conflicts.xlsx','BASE SHEET',{'ICRIS Number':'Account Number','Company Name':'Company Name'});db.commit()
        company=db.scalar(select(Company).where(Company.icris_number==token.upper()));issues=db.scalars(select(DataQualityIssue).where(DataQualityIssue.company_id==company.id,DataQualityIssue.issue_type=='company_master_name_conflict')).all();assert len(issues)==1
    finally:db.close()

def test_manual_name_is_protected_and_conflict_recorded():
    token=unique('MANUAL');db=SessionLocal()
    try:
        company=Company(icris_number=token,company_name='Human Corrected Name',normalized_name='HUMAN CORRECTED NAME',source='manual',name_source='manual',manual_override_fields=['company_name']);db.add(company);db.commit()
        batch=import_companies(db,workbook([(token,'Workbook Name')]),'master.xlsx','BASE SHEET',{'ICRIS Number':'Account Number','Company Name':'Company Name'});db.commit();db.refresh(company)
        assert company.company_name=='Human Corrected Name';assert batch.conflict_count==1;assert db.scalar(select(DataQualityIssue).where(DataQualityIssue.company_id==company.id,DataQualityIssue.issue_type=='company_master_name_conflict'))
    finally:db.close()

def test_api_mapping_preview_and_manual_entry(client):
    token=unique('API');data=workbook([(token,'API Cargo')]);files={'file':('master.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    preview=client.post('/api/v1/company-imports/preview',files=files,data={'worksheet':'BASE SHEET','mapping_json':'{"ICRIS Number":"Account Number","Company Name":"Company Name"}'})
    assert preview.status_code==200;assert preview.json()['valid_rows']==1
    manual=client.post('/api/v1/companies',json={'icris_number':f'  {token}-M  ','company_name':'Manual Cargo'});assert manual.status_code==201;assert manual.json()['icris_number']==f'{token}-M'.upper();assert 'company_name' in manual.json()['manual_override_fields']

def test_import_transaction_can_be_rolled_back():
    token=unique('ROLLBACK');db=SessionLocal()
    try:
        filename=f'{token}.xlsx';import_companies(db,workbook([(token,'Rollback Cargo')]),filename,'BASE SHEET',{'ICRIS Number':'Account Number','Company Name':'Company Name'});db.rollback()
        assert not db.scalar(select(Company).where(Company.icris_number==token));assert not db.scalar(select(CompanyImportBatch).where(CompanyImportBatch.file_name==filename))
    finally:db.close()
