import io
import uuid
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import hash_password
from app.db import SessionLocal
from app.main import SHIPMENT_EXPORT_HEADERS, app
from app.models import Shipment, User


def _sheet_rows(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers['content-type'].startswith('application/vnd.openxmlformats')
    assert 'attachment' in resp.headers['content-disposition']
    ws = load_workbook(io.BytesIO(resp.content), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == SHIPMENT_EXPORT_HEADERS
    return rows[1:]


def _seed(tag):
    db = SessionLocal()
    try:
        a = Shipment(shipment_number=f'EXP-{tag}-A', source='crm', pay_term='PP', bill_amount=Decimal('120.50'), ae_code=f'XA{tag[:4]}')
        b = Shipment(shipment_number=f'EXP-{tag}-B', source='crm', pay_term='PP', bill_amount=Decimal('80'), ae_code=f'XB{tag[:4]}')
        db.add_all([a, b]); db.commit()
        return [a.id, b.id]
    finally:
        db.close()


def _cleanup(shipment_ids=(), emails=()):
    db = SessionLocal()
    try:
        if shipment_ids: db.query(Shipment).filter(Shipment.id.in_(shipment_ids)).delete(synchronize_session=False)
        if emails: db.query(User).filter(User.email.in_(emails)).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


def test_shipments_export_matches_list_total_and_filters(client):
    tag = uuid.uuid4().hex[:8].upper()
    ids = _seed(tag)
    try:
        params = {'shipment_number': f'EXP-{tag}'}
        listed = client.get('/api/v1/shipments', params=params).json()
        rows = _sheet_rows(client.get('/api/v1/shipments/export.xlsx', params=params))
        assert len(rows) == listed['total'] == 2
        by_awb = {r[0]: r for r in rows}
        revenue_col = SHIPMENT_EXPORT_HEADERS.index('Revenue')
        assert by_awb[f'EXP-{tag}-A'][revenue_col] == 120.5  # numeric cell, not text
    finally:
        _cleanup(shipment_ids=ids)


def test_shipments_export_respects_ae_scope():
    tag = uuid.uuid4().hex[:8].upper()
    ids = _seed(tag)
    email = f'export-ae-{tag.lower()}@customer360.test'
    db = SessionLocal()
    try:
        db.add(User(email=email, display_name='Export AE', role='ae', ae_code=f'XA{tag[:4]}', password_hash=hash_password('test-password-not-real'), is_active=True))
        db.commit()
    finally:
        db.close()
    try:
        ae_client = TestClient(app)
        assert ae_client.post('/api/v1/auth/login', json={'email': email, 'password': 'test-password-not-real'}).status_code == 200
        rows = _sheet_rows(ae_client.get('/api/v1/shipments/export.xlsx', params={'shipment_number': f'EXP-{tag}'}))
        ae_col = SHIPMENT_EXPORT_HEADERS.index('AE')
        assert [r[0] for r in rows] == [f'EXP-{tag}-A']
        assert rows[0][ae_col] == f'XA{tag[:4]}'
    finally:
        _cleanup(shipment_ids=ids, emails=[email])


def test_unscoped_analytics_csv_export_is_gone(client):
    assert client.get('/api/v1/analytics/customers/export.csv').status_code == 404
