"""One-off import for the "Clean_Customer_List" workbook (Customer_Clean_1.xlsx).

Updates the company master from a hand-curated business list: segment, contact details,
and AE assignment. Distinct from company_imports.py (which targets the ICRIS master
export with its own header contract) and ae_imports.py (AE-only, different headers) —
this file's headers are its own:

    Sno | Company Name | Icris | Last Shipment | Type | Owner Name | Owner Address |
    Owner Contact No. | Email | User | Win

Decisions confirmed with the project owner 2026-08-12:
  * Type -> customer_type: SME->'SME', LA->'Large Account', IFC->'Strategic Account',
    RE/CO->'Small Customer'. A BLANK Type leaves the existing value untouched.
  * User -> AE code, applied through ae_imports._reassign_company so the existing audit
    log / override-locking / shipment propagation behaviour is reused verbatim.
  * An ICRIS not already present is created as a provisional company (same treatment
    crm_sync gives an unseen valid ICRIS).
  * Duplicate ICRIS rows: the row with the most populated fields wins; the other rows'
    company names are kept as aliases rather than discarded.
  * Owner Name has no column in the schema and is deliberately NOT imported.

Blank cells never erase existing values (docs/03-data-rules.md rule 5), and manual field
overrides are respected (rule 4).
"""
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .ae_imports import _reassign_company
from .models import ActivityLog, Company, CompanyAlias, now
from .utils import normalize_icris, normalize_name

REQUIRED_HEADERS = ('Company Name', 'Icris', 'Type', 'User')
SHEET = 'Clean_Customer_List'

TYPE_TO_SEGMENT = {
    'SME': 'SME',
    'LA': 'Large Account',
    'IFC': 'Strategic Account',
    'RE/CO': 'Small Customer',
}

# Fields this import may write, in (excel header -> Company attribute) form. Owner Name is
# intentionally absent: there is no column for it and inventing one is out of scope here.
CONTACT_MAP = {
    'Owner Contact No.': 'phone',
    'Email': 'email',
    'Owner Address': 'address',
}


class CleanListError(ValueError):
    pass


def _clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace(' ', ' ').strip()


def _populated_score(row: dict) -> int:
    """How many of the informative columns this row actually fills — used to pick a winner
    among duplicate ICRIS rows."""
    return sum(1 for k in ('Type', 'Owner Name', 'Owner Address', 'Owner Contact No.', 'Email') if row.get(k))


def parse_workbook(path_or_bytes, filename='Customer_Clean_1.xlsx'):
    if Path(filename).suffix.casefold() not in {'.xlsx', '.xlsm'}:
        raise CleanListError('Customer clean list must be an .xlsx workbook')
    workbook = load_workbook(path_or_bytes, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise CleanListError(f'Worksheet {SHEET!r} was not found (found: {workbook.sheetnames})')
    ws = workbook[SHEET]
    header_cells = next(ws.iter_rows(min_row=1, max_row=1))
    headers = [_clean(c.value) for c in header_cells]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise CleanListError(f'Missing required column(s): {missing}. Found: {headers}')

    rows = []
    for number, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        record = {headers[i]: _clean(v) for i, v in enumerate(raw) if i < len(headers)}
        if not record.get('Icris') and not record.get('Company Name'):
            continue
        record['row_number'] = number
        rows.append(record)
    return rows


def group_rows(rows):
    """Collapse duplicate ICRIS rows: richest row wins, the rest become alias candidates."""
    by_icris = {}
    for row in rows:
        key = normalize_icris(row.get('Icris', ''))
        if not key:
            continue
        by_icris.setdefault(key, []).append(row)
    grouped = []
    for icris, group in by_icris.items():
        group_sorted = sorted(group, key=_populated_score, reverse=True)
        winner = group_sorted[0]
        alias_names = [r['Company Name'] for r in group_sorted[1:] if r.get('Company Name')]
        grouped.append({'icris': icris, 'row': winner, 'alias_names': alias_names, 'duplicate_count': len(group) - 1})
    return grouped


def run_import(db: Session, rows, dry_run=True, changed_by_user_id=None):
    stats = {
        'rows_in_file': len(rows), 'unique_icris': 0, 'duplicate_rows_collapsed': 0,
        'companies_matched': 0, 'companies_created': 0,
        'segment_updated': 0, 'segment_blank_skipped': 0,
        'contact_fields_updated': 0, 'aliases_added': 0,
        'ae_reassigned': 0, 'ae_unchanged': 0, 'ae_shipments_updated': 0,
        'skipped_no_icris': 0,
    }
    details = {'created': [], 'ae_changes': [], 'segment_changes': []}

    grouped = group_rows(rows)
    stats['unique_icris'] = len(grouped)
    stats['duplicate_rows_collapsed'] = sum(g['duplicate_count'] for g in grouped)
    stats['skipped_no_icris'] = len(rows) - sum(len(g['alias_names']) + 1 for g in grouped)

    for entry in grouped:
        icris, row = entry['icris'], entry['row']
        name = row.get('Company Name') or f'Provisional {icris}'
        company = db.scalar(select(Company).where(func.upper(func.trim(Company.icris_number)) == icris))

        if not company:
            stats['companies_created'] += 1
            details['created'].append({'icris': icris, 'name': name})
            if not dry_run:
                company = Company(
                    icris_number=icris, company_name=name, normalized_name=normalize_name(name),
                    source='customer_clean_list', is_provisional=True, name_source='customer_clean_list',
                )
                db.add(company)
                db.flush()
                db.add(ActivityLog(
                    entity_type='company', entity_id=company.id, action='created',
                    description=f'Created from Clean_Customer_List import ({icris})',
                    source='customer_clean_list', metadata_json={'icris': icris},
                ))
            else:
                continue  # nothing further to compute for a company that does not exist yet
        else:
            stats['companies_matched'] += 1

        overrides = set(company.manual_override_fields or [])

        # --- segment -------------------------------------------------------------
        type_raw = (row.get('Type') or '').strip()
        if not type_raw:
            stats['segment_blank_skipped'] += 1
        else:
            segment = TYPE_TO_SEGMENT.get(type_raw.upper())
            if segment and segment != company.customer_type and 'customer_type' not in overrides:
                details['segment_changes'].append({'icris': icris, 'from': company.customer_type, 'to': segment})
                stats['segment_updated'] += 1
                if not dry_run:
                    company.customer_type = segment

        # --- contact details (blank never erases) --------------------------------
        for header, attr in CONTACT_MAP.items():
            value = (row.get(header) or '').strip()
            if not value or attr in overrides:
                continue
            if getattr(company, attr, None) != value:
                stats['contact_fields_updated'] += 1
                if not dry_run:
                    setattr(company, attr, value)

        # --- aliases from duplicate-row name variants ----------------------------
        for alias_name in entry['alias_names']:
            normalized = normalize_name(alias_name)
            if not normalized or normalized == company.normalized_name:
                continue
            if not dry_run:
                exists = db.scalar(select(CompanyAlias).where(
                    CompanyAlias.company_id == company.id,
                    CompanyAlias.normalized_alias_name == normalized,
                ))
                if exists:
                    continue
                db.add(CompanyAlias(
                    company_id=company.id, alias_name=alias_name,
                    normalized_alias_name=normalized, source='customer_clean_list',
                ))
            stats['aliases_added'] += 1

        # --- AE assignment -------------------------------------------------------
        ae_code = (row.get('User') or '').strip().upper()
        if ae_code:
            if company.assigned_ae_code == ae_code:
                stats['ae_unchanged'] += 1
            else:
                details['ae_changes'].append({'icris': icris, 'from': company.assigned_ae_code, 'to': ae_code})
                stats['ae_reassigned'] += 1
                if not dry_run:
                    stats['ae_shipments_updated'] += _reassign_company(
                        db, company, ae_code, reason='Clean_Customer_List import',
                        source='customer_clean_list', changed_by_user_id=changed_by_user_id,
                    )

    return stats, details
