"""Import for AE monthly targets ("AE_Targets_2026.xlsx", 'All_AE_Flat' sheet):

  Year | AE | MONTH | Weight | Piece | Revenue | Weight_Imp | Piece_Imp | Revenue_Imp

The AE column carries a full display label ending in the short code ("Ankit Shrestha AS"),
not the bare code alone, so it's resolved from the *last token* of that label rather than
ae_imports.py's NAME_TO_CODE map (which is keyed on first names appearing anywhere in a
different source file's AE column, not a suffix here). A label whose last token doesn't
match a known AccountExecutive.ae_code is reported as unknown rather than silently
registering a new AE — targets are numbers set by management, not an identity source.
"""
import calendar
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import AccountExecutive, AeTarget

class AeTargetWorkbookError(ValueError): pass

REQUIRED_HEADERS = ('Year', 'AE', 'MONTH')
MONTH_NAME_TO_NUMBER = {name.casefold(): i for i, name in enumerate(calendar.month_name) if name}
MONTH_NAME_TO_NUMBER.update({name.casefold(): i for i, name in enumerate(calendar.month_abbr) if name})

def clean_header(value): return ' '.join(str(value or '').replace('\xa0', ' ').split()).strip()

def extract_ae_code(raw, known_codes):
    text = str(raw or '').strip()
    if not text: return None, text
    token = text.split()[-1].upper()
    if token in known_codes: return token, text
    return None, text

def parse_month(raw):
    text = str(raw or '').strip()
    if not text: return None
    if text.isdigit():
        n = int(text)
        return n if 1 <= n <= 12 else None
    return MONTH_NAME_TO_NUMBER.get(text.casefold())

def _num(value):
    if value is None or value == '': return None
    try: return float(value)
    except (TypeError, ValueError): return None

def parse_workbook(data: bytes, filename: str, worksheet: str | None = None):
    if Path(filename).suffix.casefold() not in {'.xlsx', '.xlsm'}:
        raise AeTargetWorkbookError('AE targets file must be an .xlsx workbook')
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise AeTargetWorkbookError('The uploaded workbook is malformed or unreadable') from exc
    wanted = clean_header(worksheet or 'All_AE_Flat')
    available = {clean_header(s).casefold(): s for s in workbook.sheetnames}
    sheet_name = available.get(wanted.casefold(), workbook.sheetnames[0])
    ws = workbook[sheet_name]
    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [clean_header(h) for h in raw_headers if clean_header(h)]
    column_by_header = {}
    for index, value in enumerate(raw_headers, 1):
        header = clean_header(value)
        if header: column_by_header.setdefault(header, index)
    missing = [h for h in REQUIRED_HEADERS if h not in column_by_header]
    if missing:
        workbook.close()
        raise AeTargetWorkbookError(f"Expected columns not found: {', '.join(missing)}. Found: {', '.join(headers)}")
    col = {name: column_by_header.get(name) for name in ('Year', 'AE', 'MONTH', 'Weight', 'Piece', 'Revenue', 'Weight_Imp', 'Piece_Imp', 'Revenue_Imp')}
    rows = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2), 2):
        def cell(name):
            c = col[name]
            return row[c - 1].value if c and len(row) >= c else None
        year_raw, ae_raw, month_raw = cell('Year'), cell('AE'), cell('MONTH')
        if not any([year_raw, ae_raw, month_raw]): continue
        rows.append({
            'row_number': row_number, 'year_raw': year_raw, 'ae_raw': ae_raw, 'month_raw': month_raw,
            'weight': _num(cell('Weight')), 'piece': _num(cell('Piece')), 'revenue': _num(cell('Revenue')),
            'weight_imp': _num(cell('Weight_Imp')), 'piece_imp': _num(cell('Piece_Imp')), 'revenue_imp': _num(cell('Revenue_Imp')),
        })
    workbook.close()
    return {'file_name': Path(filename).name, 'worksheet_name': sheet_name, 'headers': headers, 'rows': rows}

def import_ae_targets(db: Session, data: bytes, filename: str, worksheet: str | None = None) -> dict:
    parsed = parse_workbook(data, filename, worksheet)
    known_codes = {a.ae_code for a in db.scalars(select(AccountExecutive)).all()}
    created = updated = 0
    unknown_ae = []
    invalid_rows = []
    for row in parsed['rows']:
        ae_code, ae_label = extract_ae_code(row['ae_raw'], known_codes)
        month = parse_month(row['month_raw'])
        try:
            year = int(row['year_raw']) if row['year_raw'] not in (None, '') else None
        except (TypeError, ValueError):
            year = None
        if not ae_code:
            unknown_ae.append({'row_number': row['row_number'], 'ae_raw': ae_label})
            continue
        if not year or not month:
            invalid_rows.append({'row_number': row['row_number'], 'year_raw': row['year_raw'], 'month_raw': row['month_raw']})
            continue
        existing = db.scalar(select(AeTarget).where(AeTarget.ae_code == ae_code, AeTarget.year == year, AeTarget.month == month))
        if not existing:
            existing = AeTarget(ae_code=ae_code, year=year, month=month, source='import')
            db.add(existing)
            created += 1
        else:
            existing.source = 'import'
            updated += 1
        existing.weight_target = row['weight']; existing.piece_target = int(row['piece']) if row['piece'] is not None else None; existing.revenue_target = row['revenue']
        existing.weight_target_import = row['weight_imp']; existing.piece_target_import = int(row['piece_imp']) if row['piece_imp'] is not None else None; existing.revenue_target_import = row['revenue_imp']
    return {
        'file_name': parsed['file_name'], 'worksheet_name': parsed['worksheet_name'],
        'total_rows': len(parsed['rows']), 'created': created, 'updated': updated,
        'unknown_ae_count': len(unknown_ae), 'unknown_ae_values': sorted({r['ae_raw'] for r in unknown_ae}),
        'invalid_row_count': len(invalid_rows), 'invalid_rows': invalid_rows[:25],
    }
