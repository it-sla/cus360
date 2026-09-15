"""Transactional, grouped import for the authoritative ICRIS company master."""
from collections import defaultdict
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import Company,CompanyAlias,CompanyImportBatch,CompanyImportRawRow,DataQualityIssue,now
from .utils import normalize_icris,normalize_name

DEFAULT_SHEET='BASE SHEET';IGNORED_SHEET='PIVOT TABLE';MAPPING_KEYS=('ICRIS Number','Company Name')
class CompanyWorkbookError(ValueError):pass
def clean_header(value):return ' '.join(str(value or '').replace('\u00a0',' ').split()).strip()
def clean_text(value):
    if value is None:return ''
    return str(value).replace('\u00a0',' ').strip(' \t\r\n\u200b\u200c\u200d\u2060\ufeff')
def cell_text(cell):
    value=cell.value
    if value is None:return ''
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        fmt=(cell.number_format or '').split(';')[0]
        if isinstance(value,int) and fmt and set(fmt)<=set('0'):return str(value).zfill(len(fmt))
        return str(value)
    return clean_text(value)
def _sheet(workbook,requested):
    wanted=clean_header(requested or DEFAULT_SHEET).casefold();available={clean_header(x).casefold():x for x in workbook.sheetnames if clean_header(x).casefold()!=IGNORED_SHEET.casefold()}
    if wanted==IGNORED_SHEET.casefold():raise CompanyWorkbookError('PIVOT TABLE is not an importable source sheet')
    if wanted not in available:raise CompanyWorkbookError(f'Worksheet {clean_header(requested or DEFAULT_SHEET)!r} was not found')
    return workbook[available[wanted]],clean_header(available[wanted]),[clean_header(x) for x in workbook.sheetnames if clean_header(x).casefold()!=IGNORED_SHEET.casefold()]
def parse_workbook(data,filename,worksheet=None,mapping=None):
    if Path(filename).suffix.casefold()!='.xlsx':raise CompanyWorkbookError('Company master must be an .xlsx workbook')
    try:workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
    except Exception as exc:raise CompanyWorkbookError('The uploaded workbook is malformed or unreadable') from exc
    ws,sheet,available=_sheet(workbook,worksheet);raw_headers=[cell.value for cell in next(ws.iter_rows(min_row=1,max_row=1))];headers=[];column_by_header={}
    for index,value in enumerate(raw_headers,1):
        header=clean_header(value)
        if header:headers.append(header);column_by_header.setdefault(header,index)
    suggested={'ICRIS Number':next((h for h in headers if h.casefold() in {'account number','icris number'}),None),'Company Name':next((h for h in headers if h.casefold()=='company name'),None)}
    selected={**suggested,**(mapping or {})}
    if any(not selected.get(k) or clean_header(selected[k]) not in column_by_header for k in MAPPING_KEYS):raise CompanyWorkbookError('Account Number and Company Name columns must be mapped')
    icris_col=column_by_header[clean_header(selected['ICRIS Number'])];name_col=column_by_header[clean_header(selected['Company Name'])];rows=[]
    for row_number,row in enumerate(ws.iter_rows(min_row=2),2):
        icris=cell_text(row[icris_col-1]) if len(row)>=icris_col else '';name=cell_text(row[name_col-1]) if len(row)>=name_col else ''
        if not icris and not name and not any(cell_text(c) for c in row):continue
        rows.append({'row_number':row_number,'source_icris':icris,'normalized_icris':normalize_icris(icris),'company_name':name})
    workbook.close();return {'file_name':Path(filename).name,'worksheet_name':sheet,'available_worksheets':available,'headers':headers,'suggested_mapping':suggested,'mapping':selected,'rows':rows}
def analyze(db:Session,parsed):
    rows=parsed['rows'];valid=[r for r in rows if r['normalized_icris'] and r['company_name']];groups=defaultdict(list)
    for row in valid:groups[row['normalized_icris']].append(row)
    existing={normalize_icris(c.icris_number):c for c in db.scalars(select(Company)).all()};duplicates=[];conflicts=[];new=present=promote=unchanged=0
    for icris,group in groups.items():
        names={normalize_name(r['company_name']):r['company_name'] for r in group};conflict=len(names)>1
        if len(group)>1:duplicates.append({'normalized_icris':icris,'row_numbers':[r['row_number'] for r in group],'names':list(names.values()),'conflict':conflict,'rows':group})
        if conflict:conflicts.append({'normalized_icris':icris,'row_numbers':[r['row_number'] for r in group],'names':list(names.values())})
        company=existing.get(icris);primary=group[0]['company_name']
        if not company:new+=1
        else:
            present+=1
            if company.is_provisional:promote+=1
            elif normalize_name(company.company_name)==normalize_name(primary):unchanged+=1
    invalid=[{'row_number':r['row_number'],'source_icris':r['source_icris'],'company_name':r['company_name'],'errors':(['Blank Account Number'] if not r['normalized_icris'] else [])+(['Blank Company Name'] if not r['company_name'] else [])} for r in rows if not r['normalized_icris'] or not r['company_name']]
    return {**{k:v for k,v in parsed.items() if k!='rows'},'total_source_rows':len(rows),'valid_rows':len(valid),'blank_icris_count':sum(not r['normalized_icris'] for r in rows),'blank_company_name_count':sum(not r['company_name'] for r in rows),'unique_normalized_icris_count':len(groups),'duplicate_icris_group_count':len(duplicates),'new_companies':new,'existing_companies':present,'provisional_companies_to_promote':promote,'unchanged_companies':unchanged,'name_conflict_count':len(conflicts),'invalid_row_count':len(invalid),'duplicate_groups':duplicates,'name_conflicts':conflicts,'invalid_rows':invalid,'preview_rows':rows[:10]}
def import_companies(db:Session,data,filename,worksheet,mapping):
    parsed=parse_workbook(data,filename,worksheet,mapping);preview=analyze(db,parsed);digest=sha256(data).hexdigest();batch=CompanyImportBatch(file_name=parsed['file_name'],file_hash=digest,worksheet_name=parsed['worksheet_name'],mapping_json=parsed['mapping'],total_rows=preview['total_source_rows'],valid_count=preview['valid_rows'],blank_icris_count=preview['blank_icris_count'],blank_name_count=preview['blank_company_name_count'],unique_icris_count=preview['unique_normalized_icris_count'],duplicate_group_count=preview['duplicate_icris_group_count'],conflict_count=preview['name_conflict_count'],rejected_count=preview['invalid_row_count'],status='processing');db.add(batch);db.flush();groups=defaultdict(list)
    for row in parsed['rows']:
        if row['normalized_icris'] and row['company_name']:groups[row['normalized_icris']].append(row)
        else:batch.failed_count+=1;db.add(CompanyImportRawRow(import_batch_id=batch.id,row_number=row['row_number'],raw_data_json=row,processing_status='rejected',error_message='Blank Account Number or Company Name'))
    existing={normalize_icris(c.icris_number):c for c in db.scalars(select(Company)).all()}
    for icris,group in groups.items():
        distinct={normalize_name(r['company_name']):r['company_name'] for r in group};primary=group[0]['company_name'];company=existing.get(icris);status='unchanged'
        if not company:company=Company(icris_number=icris,company_name=primary,normalized_name=normalize_name(primary),source='upload',name_source='company_master',is_provisional=False,last_company_import_at=now(),last_company_import_batch_id=batch.id);db.add(company);db.flush();existing[icris]=company;batch.created_count+=1;status='created'
        else:
            previous=company.company_name;overrides=set(company.manual_override_fields or [])
            if company.is_provisional:
                company.is_provisional=False;batch.promoted_count+=1;status='promoted'
                for oi in db.scalars(select(DataQualityIssue).where(DataQualityIssue.company_id==company.id,DataQualityIssue.issue_type=='crm_icris_not_in_master',DataQualityIssue.status=='open')).all():
                    oi.status='resolved';oi.resolved_at=now();oi.resolved_by='company_master_import'
            if 'company_name' not in overrides and normalize_name(previous)!=normalize_name(primary):company.company_name=primary;company.normalized_name=normalize_name(primary);batch.updated_count+=1;status='updated' if status=='unchanged' else status
            elif status=='unchanged':batch.unchanged_count+=1
            if previous and normalize_name(previous)!=company.normalized_name:distinct.setdefault(normalize_name(previous),previous)
            company.source='upload';company.name_source='company_master';company.last_company_import_at=now();company.last_company_import_batch_id=batch.id
        for normalized,name in distinct.items():
            if normalized and normalized!=company.normalized_name and not db.scalar(select(CompanyAlias).where(CompanyAlias.company_id==company.id,CompanyAlias.normalized_alias_name==normalized)):
                db.add(CompanyAlias(company_id=company.id,alias_name=name,normalized_alias_name=normalized,source='company_master'));batch.alias_added_count+=1
        source_conflict=len({normalize_name(r['company_name']) for r in group})>1
        manual_conflict=bool(company.manual_override_fields and 'company_name' in company.manual_override_fields and normalize_name(company.company_name)!=normalize_name(primary))
        if manual_conflict and not source_conflict:batch.conflict_count+=1
        if source_conflict or manual_conflict:
            details={'import_batch_id':str(batch.id),'row_numbers':[r['row_number'] for r in group],'distinct_names':list(distinct.values())}
            issue=db.scalar(select(DataQualityIssue).where(DataQualityIssue.issue_type=='company_master_name_conflict',DataQualityIssue.company_id==company.id,DataQualityIssue.status.in_(('open','acknowledged'))).order_by(DataQualityIssue.first_seen_at))
            if issue:issue.source_company_name=primary;issue.details_json=details;issue.last_seen_at=now()
            else:db.add(DataQualityIssue(issue_type='company_master_name_conflict',severity='warning',company_id=company.id,source_icris_number=icris,source_company_name=primary,details_json=details,status='open'))
        for row in group:db.add(CompanyImportRawRow(import_batch_id=batch.id,row_number=row['row_number'],raw_data_json=row,processing_status=status,error_message=None))
    batch.status='completed_with_conflicts' if batch.conflict_count else ('completed_with_errors' if batch.failed_count else 'completed');batch.completed_at=now();return batch
